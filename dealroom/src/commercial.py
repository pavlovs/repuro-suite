"""
DR-COMMERCIAL: Commercial data extraction pipeline.
Phase 1: Customer list extraction + concentration KPIs.

Reads customer lists (Cat RAB nach Umsatz, Octopus Kundenumsatz), classifies
columns via Claude CLI, loads to deal_customers, derives concentration KPIs
to deal_commercial, and cross-checks against deal_financials.

Called by DEALROOM.py CLI commands.

ANTI-PATTERNS (enforced in this module):
  AP-1 ASSUME-COLUMN-POSITIONS: Never hardcode column positions. Use AI
       classification + human confirm.
  AP-2 IGNORE-GERMAN-NUMBERS: Always use _parse_german_number() for all
       numeric values from xlsx.
  AP-3 CUSTOMER-NAME-WITHOUT-ID-TRACKING: Extract customer_id when available
       (Cat: ID in parens after name).
  AP-4 SKIP-REVIEW-FOR-LARGE-DATASETS: Always cross-check customer revenue
       total vs Gesamtleistung.
  AP-5 DERIVE-KPIS-BEFORE-APPROVAL: Mark derived KPIs with
       confidence='computed:preliminary' when source is_authoritative=0.
"""

from __future__ import annotations

import json
import logging
import os
import re
import shutil
import subprocess
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl


log = logging.getLogger(__name__)

# ---- Claude CLI ----------------------------------------------------------

_CLAUDE_CMD = shutil.which("claude") or r"C:\Users\X1\AppData\Roaming\npm\claude.cmd"


def _call_claude(prompt: str) -> str:
    """Call Claude via CLI (OAuth). Never uses Anthropic SDK."""
    env = os.environ.copy()
    env.pop("ANTHROPIC_API_KEY", None)
    for attempt in range(2):
        try:
            result = subprocess.run(
                [_CLAUDE_CMD, "-p", "--output-format", "text"],
                input=prompt,
                capture_output=True,
                text=True,
                encoding="utf-8",
                timeout=120,
                env=env,
            )
            if result.returncode == 0:
                return result.stdout.strip()
            if attempt == 0:
                continue
            raise RuntimeError(f"Claude CLI failed: {result.stderr[:500]}")
        except subprocess.TimeoutExpired:
            if attempt == 0:
                continue
            raise
    return ""


# ---- Utilities -----------------------------------------------------------


def _parse_german_number(val: Any) -> float | None:
    """Parse German number format.

    Handles: '4.240.640,96', '(1.234,56)', '-1234', 1234.56 (already float).
    Returns None for non-parseable values.
    """
    if isinstance(val, (int, float)):
        return float(val)
    if not isinstance(val, str):
        return None
    s = val.strip()
    if not s:
        return None
    # Strip currency symbols and whitespace
    s = s.replace("€", "").replace("EUR", "").replace("eur", "").strip()
    if not s:
        return None
    # Handle parenthesised negatives: (1.234,56) -> -1234.56
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative = True
        s = s[1:-1].strip()
    elif s.startswith("-"):
        negative = True
        s = s[1:].strip()
    # German format: dot = thousands, comma = decimal
    # Detect format: if there is a comma, treat it as decimal separator
    if "," in s:
        s = s.replace(".", "")  # remove thousands separators
        s = s.replace(",", ".")  # decimal comma -> decimal point
    else:
        # No comma: dots could be thousands separators (e.g. "4.240.640")
        # or already a decimal point (e.g. "1234.56").
        # Heuristic: multiple dots = thousands separators
        if s.count(".") > 1:
            s = s.replace(".", "")
    try:
        result = float(s)
        return -result if negative else result
    except ValueError:
        return None


def _col_letter_to_index(letter: str) -> int:
    """Convert Excel column letter to 0-based index. A=0, B=1, ..., Z=25, AA=26."""
    letter = letter.upper().strip()
    idx = 0
    for ch in letter:
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    return idx - 1


def _eur_to_eurk(v: float | None) -> float | None:
    """Convert EUR to EUR K (divide by 1000)."""
    if v is None:
        return None
    return round(v / 1000.0, 2)


# ---- File reading --------------------------------------------------------


def _read_xlsx_preview(path: Path, max_rows: int = 25) -> list[list]:
    """Read first N rows from xlsx for AI classification preview.

    Returns list of lists (row values as strings). Uses openpyxl read_only.
    """
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        raise ValueError(f"Cannot open {path.name}: {e}") from e

    result = []
    ws = wb.worksheets[0]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_rows:
            break
        # Stringify cells for the AI preview
        cells = []
        for cell in row:
            if cell is None:
                cells.append("")
            else:
                cells.append(str(cell))
        result.append(cells)
    wb.close()
    return result


def _read_xlsx_full(path: Path, sheet_index: int = 0) -> list[tuple]:
    """Read all rows from an xlsx sheet. Returns list of row tuples (raw values)."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        raise ValueError(f"Cannot open {path.name}: {e}") from e

    ws = (
        wb.worksheets[sheet_index]
        if sheet_index < len(wb.worksheets)
        else wb.worksheets[0]
    )
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(tuple(row))
    wb.close()
    return rows


# ---- AI Classification --------------------------------------------------


_CLASSIFY_PROMPT_TEMPLATE = """You are a data analyst classifying an Excel file for M&A due diligence.

FILE NAME: {file_name}
SHEET NAME: {sheet_name}
PREVIEW (first {n_rows} rows, tab-separated):
{preview_text}

Classify this file and map its columns.

RULES:
1. Identify file_type: one of customer_list, product_split, backlog, employee_list, supplier_list
2. Identify target_table: deal_customers, deal_products, deal_backlog, deal_employees, deal_suppliers
3. Map each column letter (A, B, C...) to a target field with a transform
4. Identify header_row (1-indexed) and data_start_row (1-indexed)
5. Count footer/total rows to skip (rows at bottom with totals or empty)
6. Detect fiscal_year from header text or filename if present

Available target fields for deal_customers:
  customer_name, customer_id, revenue_k (divide raw EUR by 1000), revenue_raw (keep raw EUR),
  revenue_pct, rank, cohort, customer_type, specialty, notes

Available transforms:
  text - pass through as text
  extract_name_before_paren - "Mueller GmbH (12345)" -> "Mueller GmbH"
  extract_id_in_paren - "Mueller GmbH (12345)" -> "12345"
  german_number - parse German number format (dot=thousands, comma=decimal)
  german_number_div_1000 - parse German number, divide by 1000
  abc_letter - map to A/B/C cohort
  percentage - parse as percentage (already 0-100)
  integer - parse as integer
  skip - ignore this column

Return JSON wrapped in ```json fences:
```json
{{
  "file_type": "customer_list",
  "target_table": "deal_customers",
  "fiscal_year": 2024,
  "sheets": {{
    "{sheet_name}": {{
      "header_row": 1,
      "data_start_row": 2,
      "skip_footer_rows": 1,
      "column_map": [
        {{"col": "A", "target": "rank", "transform": "integer"}},
        {{"col": "B", "target": "customer_name", "transform": "extract_name_before_paren"}},
        {{"col": "B", "target": "customer_id", "transform": "extract_id_in_paren"}},
        {{"col": "C", "target": "revenue_raw", "transform": "german_number"}},
        {{"col": "D", "target": "revenue_pct", "transform": "percentage"}},
        {{"col": "E", "target": "cohort", "transform": "abc_letter"}}
      ]
    }}
  }}
}}
```

Only return the JSON block, no other text."""


def classify_file(conn, doc_id: str) -> dict:
    """Propose column mapping for a commercial file via Claude CLI.

    Reads the file header, sends to Claude with classification prompt.
    Returns mapping dict with file_type, target_table, fiscal_year, sheets.
    """
    row = conn.execute(
        "SELECT file_path, file_name FROM deal_documents WHERE id = ?", (doc_id,)
    ).fetchone()
    if not row:
        raise ValueError(f"Document not found: {doc_id}")

    path = Path(row["file_path"])
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    # Read preview rows
    preview_rows = _read_xlsx_preview(path, max_rows=25)
    if not preview_rows:
        raise ValueError(f"No data in {path.name}")

    # Get sheet name
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet_name = wb.sheetnames[0]
        wb.close()
    except Exception:
        sheet_name = "Sheet1"

    # Build tab-separated preview
    preview_lines = []
    for i, cells in enumerate(preview_rows, 1):
        preview_lines.append(f"Row {i}: " + "\t".join(cells))
    preview_text = "\n".join(preview_lines)

    prompt = _CLASSIFY_PROMPT_TEMPLATE.format(
        file_name=row["file_name"],
        sheet_name=sheet_name,
        n_rows=len(preview_rows),
        preview_text=preview_text,
    )

    response = _call_claude(prompt)
    if not response:
        raise RuntimeError("Claude CLI returned empty response for classification")

    # Parse JSON from response (expect ```json ... ``` fences)
    json_match = re.search(r"```json\s*\n?(.*?)\n?\s*```", response, re.DOTALL)
    if json_match:
        json_str = json_match.group(1)
    else:
        # Try parsing the whole response as JSON
        json_str = response

    try:
        mapping = json.loads(json_str)
    except json.JSONDecodeError as e:
        raise ValueError(
            f"Could not parse classification response as JSON: {e}\nResponse: {response[:500]}"
        ) from e

    return mapping


# ---- Mapping persistence -------------------------------------------------


def confirm_mapping(conn, doc_id: str, mapping: dict) -> None:
    """Save confirmed mapping to deal_documents.extraction_config."""
    conn.execute(
        "UPDATE deal_documents SET extraction_config = ? WHERE id = ?",
        (json.dumps(mapping, ensure_ascii=False), doc_id),
    )
    conn.commit()


def get_mapping(conn, doc_id: str) -> dict | None:
    """Load saved mapping from deal_documents.extraction_config."""
    row = conn.execute(
        "SELECT extraction_config FROM deal_documents WHERE id = ?", (doc_id,)
    ).fetchone()
    if row and row[0]:
        return json.loads(row[0])
    return None


# ---- Transform application -----------------------------------------------


def _apply_transforms(
    row_values: tuple,
    column_map: list[dict],
    header_row_values: tuple | None = None,
) -> dict:
    """Apply column mapping transforms to a single data row.

    column_map format: [{"col": "A", "target": "customer_name", "transform": "extract_name_before_paren"}, ...]
    Returns dict of {target_field: transformed_value}.
    """
    result: dict[str, Any] = {}

    for mapping in column_map:
        col_letter = mapping["col"]
        target = mapping["target"]
        transform = mapping.get("transform", "text")

        if transform == "skip":
            continue

        col_idx = _col_letter_to_index(col_letter)
        if col_idx >= len(row_values):
            continue

        raw = row_values[col_idx]
        if raw is None:
            continue

        raw_str = str(raw).strip() if not isinstance(raw, (int, float)) else raw

        if transform == "text":
            result[target] = str(raw).strip() if raw else None

        elif transform == "extract_name_before_paren":
            s = str(raw).strip()
            # "Mueller GmbH (12345)" -> "Mueller GmbH"
            m = re.match(r"^(.+?)\s*\(", s)
            result[target] = m.group(1).strip() if m else s

        elif transform == "extract_id_in_paren":
            s = str(raw).strip()
            m = re.search(r"\(([^)]+)\)", s)
            result[target] = m.group(1).strip() if m else None

        elif transform == "german_number":
            result[target] = _parse_german_number(raw)

        elif transform == "german_number_div_1000":
            parsed = _parse_german_number(raw)
            result[target] = _eur_to_eurk(parsed)

        elif transform == "abc_letter":
            s = str(raw).strip().upper()
            if s in ("A", "B", "C"):
                result[target] = s
            else:
                result[target] = s[:1] if s else None

        elif transform == "percentage":
            parsed = _parse_german_number(raw)
            result[target] = parsed

        elif transform == "integer":
            parsed = _parse_german_number(raw)
            result[target] = int(parsed) if parsed is not None else None

        else:
            log.warning("Unknown transform %r for column %s", transform, col_letter)
            result[target] = str(raw).strip() if raw else None

    return result


# ---- Data loading --------------------------------------------------------


def _load_customers(
    conn,
    domain: str,
    rows: list[dict],
    source_file_id: str,
    fiscal_year: int,
    source_name: str,
    dry_run: bool,
) -> int:
    """Write customer rows to deal_customers. Idempotent via source_file_id."""
    if dry_run:
        return len(rows)

    # Idempotent: clear previous load from this source file
    conn.execute(
        "DELETE FROM deal_customers WHERE source_file_id = ?", (source_file_id,)
    )

    now = datetime.now(timezone.utc).isoformat()
    written = 0

    for row_data in rows:
        customer_name = row_data.get("customer_name")
        if not customer_name or not customer_name.strip():
            continue

        # Compute revenue_k: prefer pre-computed, else derive from revenue_raw
        revenue_k = row_data.get("revenue_k")
        if revenue_k is None:
            revenue_raw = row_data.get("revenue_raw")
            if revenue_raw is not None:
                revenue_k = _eur_to_eurk(revenue_raw)

        conn.execute(
            """INSERT OR REPLACE INTO deal_customers
               (id, domain, customer_name, customer_id, fiscal_year,
                revenue_k, revenue_pct, rank, cohort, customer_type,
                specialty, notes, source, source_file_id, extracted_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                str(uuid.uuid4()),
                domain,
                customer_name.strip(),
                row_data.get("customer_id"),
                fiscal_year,
                revenue_k,
                row_data.get("revenue_pct"),
                row_data.get("rank"),
                row_data.get("cohort"),
                row_data.get("customer_type"),
                row_data.get("specialty"),
                row_data.get("notes"),
                source_name,
                source_file_id,
                now,
            ),
        )
        written += 1

    conn.commit()
    return written


def _load_products(
    conn,
    domain: str,
    rows: list[dict],
    source_file_id: str,
    fiscal_year: int,
    source_name: str,
    dry_run: bool,
) -> int:
    """Write product rows to deal_products. Handles multi-year pivot."""
    # Resolve alias: AI classifier may map to "product_group" instead of "product_name"
    _PRODUCT_NAME_ALIASES = ("product_group",)

    if dry_run:
        count = 0
        for row_data in rows:
            pname = row_data.get("product_name")
            if not pname:
                for alias in _PRODUCT_NAME_ALIASES:
                    pname = row_data.get(alias)
                    if pname:
                        break
            if not pname:
                continue
            year_cols = [
                k for k in row_data if k.startswith("revenue_") and k.endswith("_raw")
            ]
            count += max(len(year_cols), 1)
        return count

    conn.execute(
        "DELETE FROM deal_products WHERE source_file_id = ?", (source_file_id,)
    )

    now = datetime.now(timezone.utc).isoformat()
    written = 0

    for row_data in rows:
        product_name = row_data.get("product_name")
        if not product_name:
            for alias in _PRODUCT_NAME_ALIASES:
                product_name = row_data.get(alias)
                if product_name:
                    break
        if not product_name or not product_name.strip():
            continue

        year_cols = {
            k: v
            for k, v in row_data.items()
            if k.startswith("revenue_") and k.endswith("_raw") and v is not None
        }

        if year_cols:
            for col_name, raw_val in year_cols.items():
                yr = _extract_year_from_field(col_name, fiscal_year)
                revenue_k = _eur_to_eurk(raw_val)
                conn.execute(
                    """INSERT OR REPLACE INTO deal_products
                       (id, domain, product_name, category, subcategory, fiscal_year,
                        revenue_k, cost_k, gross_profit_k, margin_pct, revenue_share_pct,
                        units_sold, is_recurring, notes, source, source_file_id,
                        confidence, is_authoritative, extracted_at)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        str(uuid.uuid4()),
                        domain,
                        product_name.strip(),
                        row_data.get("category"),
                        row_data.get("subcategory"),
                        yr,
                        revenue_k,
                        _eur_to_eurk(row_data.get("cost_raw"))
                        if row_data.get("cost_raw")
                        else None,
                        _eur_to_eurk(row_data.get("gross_profit_raw"))
                        if row_data.get("gross_profit_raw")
                        else None,
                        row_data.get("margin_pct"),
                        None,
                        row_data.get("units_sold"),
                        row_data.get("is_recurring", 0),
                        row_data.get("notes"),
                        source_name,
                        source_file_id,
                        "stated",
                        0,
                        now,
                    ),
                )
                written += 1
        else:
            revenue_k = row_data.get("revenue_k")
            if revenue_k is None and row_data.get("revenue_raw") is not None:
                revenue_k = _eur_to_eurk(row_data["revenue_raw"])
            conn.execute(
                """INSERT OR REPLACE INTO deal_products
                   (id, domain, product_name, fiscal_year, revenue_k,
                    source, source_file_id, confidence, is_authoritative, extracted_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    str(uuid.uuid4()),
                    domain,
                    product_name.strip(),
                    fiscal_year,
                    revenue_k,
                    source_name,
                    source_file_id,
                    "stated",
                    0,
                    now,
                ),
            )
            written += 1

    conn.commit()
    return written


def _extract_year_from_field(field_name: str, fallback_year: int) -> int:
    """Extract year from field names like 'revenue_2023_raw'."""
    m = re.search(r"(\d{4})", field_name)
    return int(m.group(1)) if m else fallback_year


def _load_backlog(
    conn,
    domain: str,
    rows: list[dict],
    source_file_id: str,
    fiscal_year: int,
    source_name: str,
    dry_run: bool,
) -> int:
    """Write backlog rows to deal_backlog. Idempotent via source_file_id."""
    if dry_run:
        return sum(
            1 for r in rows if r.get("customer_name") or r.get("project_description")
        )

    conn.execute("DELETE FROM deal_backlog WHERE source_file_id = ?", (source_file_id,))

    now = datetime.now(timezone.utc).isoformat()
    written = 0

    for row_data in rows:
        if not row_data.get("customer_name") and not row_data.get(
            "project_description"
        ):
            continue

        order_value_k = row_data.get("order_value_k")
        if order_value_k is None:
            raw = row_data.get("order_value_raw") or row_data.get("order_value")
            if raw is not None:
                order_value_k = _eur_to_eurk(raw)

        conn.execute(
            """INSERT INTO deal_backlog
               (id, domain, customer_name, project_description, location,
                execution_year, order_value_k, margin_pct, status,
                fiscal_year, source, source_file_id, confidence, is_authoritative, extracted_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                str(uuid.uuid4()),
                domain,
                (row_data.get("customer_name") or "").strip() or None,
                (row_data.get("project_description") or "").strip() or None,
                (row_data.get("location") or "").strip() or None,
                (row_data.get("execution_year") or "").strip()
                if row_data.get("execution_year")
                else None,
                order_value_k,
                row_data.get("margin_pct"),
                (row_data.get("status") or "").strip() or None,
                fiscal_year,
                source_name,
                source_file_id,
                "stated",
                0,
                now,
            ),
        )
        written += 1

    conn.commit()
    return written


def load_file(conn, doc_id: str, dry_run: bool = False) -> dict:
    """Load commercial data from file using confirmed mapping.

    Returns: {rows_written: int, rows_skipped: int, table: str, fiscal_year: int}
    """
    mapping = get_mapping(conn, doc_id)
    if not mapping:
        raise ValueError(
            f"No extraction_config for doc {doc_id}. Run classify_file first."
        )

    doc = conn.execute(
        "SELECT file_path, file_name, domain, code_name FROM deal_documents WHERE id = ?",
        (doc_id,),
    ).fetchone()
    if not doc:
        raise ValueError(f"Document not found: {doc_id}")

    path = Path(doc["file_path"])
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    domain = doc["domain"]
    target_table = mapping.get("target_table", "deal_customers")
    fiscal_year = mapping.get("fiscal_year")

    # Read full data
    all_rows = _read_xlsx_full(path)
    if not all_rows:
        return {
            "rows_written": 0,
            "rows_skipped": 0,
            "table": target_table,
            "fiscal_year": fiscal_year,
        }

    # Get sheet config (use first sheet config)
    sheets = mapping.get("sheets", {})
    if not sheets:
        raise ValueError("Mapping has no sheet configuration")
    sheet_config = next(iter(sheets.values()))

    header_row_idx = sheet_config.get("header_row", 1) - 1  # 0-indexed
    data_start_idx = sheet_config.get("data_start_row", 2) - 1
    skip_footer = sheet_config.get("skip_footer_rows", 0)
    column_map = sheet_config.get("column_map", [])

    if not column_map:
        raise ValueError("Column map is empty")

    # Extract header row for reference
    header_values = all_rows[header_row_idx] if header_row_idx < len(all_rows) else None

    # Slice data rows (exclude header and footer)
    end_idx = len(all_rows) - skip_footer if skip_footer > 0 else len(all_rows)
    data_rows = all_rows[data_start_idx:end_idx]

    # Apply transforms to each row
    transformed = []
    skipped = 0
    for row_values in data_rows:
        # Skip completely empty rows
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in row_values):
            skipped += 1
            continue
        row_dict = _apply_transforms(row_values, column_map, header_values)
        if row_dict:
            transformed.append(row_dict)
        else:
            skipped += 1

    # Route to target table loader
    loader_args = (
        conn,
        domain,
        transformed,
        doc_id,
        fiscal_year,
        doc["file_name"],
        dry_run,
    )
    if target_table == "deal_customers":
        n = _load_customers(*loader_args)
    elif target_table == "deal_products":
        n = _load_products(*loader_args)
    elif target_table == "deal_backlog":
        n = _load_backlog(*loader_args)
    else:
        raise NotImplementedError(f"Loader for {target_table} not implemented yet")

    return {
        "rows_written": n,
        "rows_skipped": skipped,
        "table": target_table,
        "fiscal_year": fiscal_year,
    }


# ---- KPI Derivation -----------------------------------------------------


def _compute_customer_kpis(customers: list[dict]) -> list[dict]:
    """Compute concentration + cohort KPIs from customer rows.

    Returns list of dicts ready for deal_commercial insertion:
      [{metric, value_num, unit, category}, ...]
    """
    if not customers:
        return []

    kpis: list[dict] = []

    # Sort by revenue descending for concentration calc
    sorted_custs = sorted(
        customers, key=lambda c: c.get("revenue_k") or 0, reverse=True
    )

    total_revenue = sum(c.get("revenue_k") or 0 for c in sorted_custs)
    n_customers = len(sorted_custs)

    kpis.append(
        {
            "metric": "customer_count",
            "value_num": n_customers,
            "unit": "#",
            "category": "customers",
        }
    )

    if total_revenue > 0:
        # Revenue shares
        shares = [(c.get("revenue_k") or 0) / total_revenue for c in sorted_custs]

        # Top N concentration
        if n_customers >= 1:
            kpis.append(
                {
                    "metric": "top1_customer_pct",
                    "value_num": round(shares[0] * 100, 2),
                    "unit": "%",
                    "category": "customers",
                }
            )
        for n, label in [(3, "top3"), (5, "top5"), (10, "top10"), (20, "top20")]:
            if n_customers >= 1:
                top_n_pct = sum(shares[: min(n, n_customers)]) * 100
                kpis.append(
                    {
                        "metric": f"{label}_customer_pct",
                        "value_num": round(top_n_pct, 2),
                        "unit": "%",
                        "category": "customers",
                    }
                )

        # Herfindahl index: sum of squared shares
        hhi = sum(s**2 for s in shares)
        kpis.append(
            {
                "metric": "herfindahl_index",
                "value_num": round(hhi, 6),
                "unit": "index",
                "category": "customers",
            }
        )

        # Revenue per customer
        kpis.append(
            {
                "metric": "revenue_per_customer_k",
                "value_num": round(total_revenue / n_customers, 2),
                "unit": "K EUR",
                "category": "customers",
            }
        )

    # Cohort KPIs (if ABC classification available)
    cohorts_present = {c.get("cohort") for c in sorted_custs if c.get("cohort")}
    if cohorts_present:
        for cohort_letter in ("A", "B", "C"):
            cohort_custs = [c for c in sorted_custs if c.get("cohort") == cohort_letter]
            cohort_count = len(cohort_custs)
            kpis.append(
                {
                    "metric": f"cohort_{cohort_letter.lower()}_count",
                    "value_num": cohort_count,
                    "unit": "#",
                    "category": "customers",
                }
            )
            if total_revenue > 0:
                cohort_rev = sum(c.get("revenue_k") or 0 for c in cohort_custs)
                kpis.append(
                    {
                        "metric": f"cohort_{cohort_letter.lower()}_pct",
                        "value_num": round(cohort_rev / total_revenue * 100, 2),
                        "unit": "%",
                        "category": "customers",
                    }
                )

    return kpis


def derive_kpis(
    conn, domain: str, fiscal_year: int | None = None, dry_run: bool = False
) -> dict:
    """Compute customer concentration KPIs and store in deal_commercial.

    Returns: {kpis_written: int, fiscal_years: list[int]}
    """
    # Determine which fiscal years have customer data
    if fiscal_year is not None:
        years = [fiscal_year]
    else:
        rows = conn.execute(
            "SELECT DISTINCT fiscal_year FROM deal_customers WHERE domain = ? ORDER BY fiscal_year",
            (domain,),
        ).fetchall()
        years = [r["fiscal_year"] for r in rows if r["fiscal_year"]]

    if not years:
        return {"kpis_written": 0, "fiscal_years": []}

    total_written = 0
    now = datetime.now(timezone.utc).isoformat()

    # Check if source customer data is authoritative
    # AP-5: If source data is not authoritative, mark KPIs as preliminary
    auth_check = conn.execute(
        "SELECT MIN(COALESCE(is_authoritative, 0)) as min_auth "
        "FROM deal_customers WHERE domain = ?",
        (domain,),
    ).fetchone()
    # If any customer row is not authoritative, KPIs are preliminary
    is_auth = bool(auth_check and auth_check["min_auth"])

    for yr in years:
        # Query customer data for this year
        cust_rows = conn.execute(
            """SELECT customer_name, customer_id, revenue_k, revenue_pct, cohort
               FROM deal_customers
               WHERE domain = ? AND fiscal_year = ?
               ORDER BY revenue_k DESC""",
            (domain, yr),
        ).fetchall()

        if not cust_rows:
            continue

        customers = [dict(r) for r in cust_rows]

        # Compute revenue_pct if not provided (from raw revenue_k)
        total_rev = sum(c.get("revenue_k") or 0 for c in customers)
        if total_rev > 0:
            for c in customers:
                if c.get("revenue_pct") is None and c.get("revenue_k") is not None:
                    c["revenue_pct"] = round(c["revenue_k"] / total_rev * 100, 2)

        kpis = _compute_customer_kpis(customers)

        if dry_run:
            total_written += len(kpis)
            continue

        # Delete previous computed KPIs for this domain/year/category=customers
        conn.execute(
            """DELETE FROM deal_commercial
               WHERE domain = ? AND fiscal_year = ? AND category = 'customers'
               AND confidence LIKE 'computed:%'""",
            (domain, yr),
        )

        # AP-5: Mark confidence based on source authoritativeness
        confidence = "computed:authoritative" if is_auth else "computed:preliminary"

        for kpi in kpis:
            conn.execute(
                """INSERT INTO deal_commercial
                   (id, domain, category, metric, fiscal_year, value_num, unit,
                    source, confidence, is_authoritative, extracted_at)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    str(uuid.uuid4()),
                    domain,
                    kpi["category"],
                    kpi["metric"],
                    yr,
                    kpi["value_num"],
                    kpi["unit"],
                    "commercial.derive_kpis",
                    confidence,
                    0,  # is_authoritative=0 for computed KPIs
                    now,
                ),
            )
            total_written += 1

        conn.commit()

    return {"kpis_written": total_written, "fiscal_years": years}


# ---- Cross-check --------------------------------------------------------


def cross_check_revenue(conn, domain: str, fiscal_year: int) -> dict | None:
    """Compare customer revenue total vs Gesamtleistung from deal_financials.

    Returns: {customer_total_k, gesamtleistung_k, delta_pct, status}
    Status: 'ok' if <2%, 'warning' if 2-5%, 'mismatch' if >5%.
    """
    # Sum customer revenue for this year
    cust_row = conn.execute(
        "SELECT SUM(revenue_k) as total FROM deal_customers WHERE domain = ? AND fiscal_year = ?",
        (domain, fiscal_year),
    ).fetchone()

    if not cust_row or cust_row["total"] is None:
        return None

    customer_total_k = cust_row["total"]

    # Get Gesamtleistung from deal_financials
    fin_row = conn.execute(
        """SELECT value_k FROM deal_financials
           WHERE domain = ? AND line_item = 'gesamtleistung'
           AND fiscal_year = ? AND period_type = 'annual'
           ORDER BY is_authoritative DESC, extracted_at DESC
           LIMIT 1""",
        (domain, fiscal_year),
    ).fetchone()

    if not fin_row or fin_row["value_k"] is None:
        # Try revenue if gesamtleistung not available
        fin_row = conn.execute(
            """SELECT value_k FROM deal_financials
               WHERE domain = ? AND line_item = 'revenue'
               AND fiscal_year = ? AND period_type = 'annual'
               ORDER BY is_authoritative DESC, extracted_at DESC
               LIMIT 1""",
            (domain, fiscal_year),
        ).fetchone()

    if not fin_row or fin_row["value_k"] is None:
        return {
            "customer_total_k": round(customer_total_k, 2),
            "gesamtleistung_k": None,
            "delta_pct": None,
            "status": "no_financial_data",
        }

    gesamtleistung_k = fin_row["value_k"]

    if abs(gesamtleistung_k) < 0.01:
        return {
            "customer_total_k": round(customer_total_k, 2),
            "gesamtleistung_k": round(gesamtleistung_k, 2),
            "delta_pct": None,
            "status": "zero_gesamtleistung",
        }

    delta_pct = abs(customer_total_k - gesamtleistung_k) / abs(gesamtleistung_k) * 100

    if delta_pct < 2:
        status = "ok"
    elif delta_pct < 5:
        status = "warning"
    else:
        status = "mismatch"

    return {
        "customer_total_k": round(customer_total_k, 2),
        "gesamtleistung_k": round(gesamtleistung_k, 2),
        "delta_pct": round(delta_pct, 2),
        "status": status,
    }


# ---- Review summary + approval ------------------------------------------


def review_summary(conn, domain: str) -> dict:
    """Build review summary for Roman.

    Returns: {
        tables: {deal_customers: {count, fiscal_years, is_authoritative}},
        top_customers: [{name, revenue_k, pct}],  # top 10 latest year
        cross_checks: [{fiscal_year, status, delta_pct}],
        kpis: [{metric, value, unit, fiscal_year}]
    }
    """
    summary: dict[str, Any] = {
        "tables": {},
        "top_customers": [],
        "cross_checks": [],
        "kpis": [],
    }

    # Customer table stats
    cust_stats = conn.execute(
        """SELECT COUNT(*) as cnt,
                  GROUP_CONCAT(DISTINCT fiscal_year) as years
           FROM deal_customers WHERE domain = ?""",
        (domain,),
    ).fetchone()

    if cust_stats and cust_stats["cnt"]:
        years_str = cust_stats["years"] or ""
        fiscal_years = sorted(int(y) for y in years_str.split(",") if y.strip())
        summary["tables"]["deal_customers"] = {
            "count": cust_stats["cnt"],
            "fiscal_years": fiscal_years,
            "is_authoritative": False,
        }

        # Top 10 customers for latest year
        latest_year = max(fiscal_years) if fiscal_years else None
        if latest_year:
            top10 = conn.execute(
                """SELECT customer_name, revenue_k, revenue_pct
                   FROM deal_customers
                   WHERE domain = ? AND fiscal_year = ?
                   ORDER BY revenue_k DESC
                   LIMIT 10""",
                (domain, latest_year),
            ).fetchall()
            summary["top_customers"] = [
                {
                    "name": r["customer_name"],
                    "revenue_k": r["revenue_k"],
                    "pct": r["revenue_pct"],
                }
                for r in top10
            ]

            # Cross-checks per fiscal year
            for yr in fiscal_years:
                check = cross_check_revenue(conn, domain, yr)
                if check:
                    summary["cross_checks"].append(
                        {
                            "fiscal_year": yr,
                            "status": check["status"],
                            "delta_pct": check.get("delta_pct"),
                            "customer_total_k": check["customer_total_k"],
                            "gesamtleistung_k": check.get("gesamtleistung_k"),
                        }
                    )

    # KPIs from deal_commercial
    kpi_rows = conn.execute(
        """SELECT metric, value_num, unit, fiscal_year
           FROM deal_commercial
           WHERE domain = ? AND category = 'customers'
           ORDER BY fiscal_year DESC, metric""",
        (domain,),
    ).fetchall()
    summary["kpis"] = [
        {
            "metric": r["metric"],
            "value": r["value_num"],
            "unit": r["unit"],
            "fiscal_year": r["fiscal_year"],
        }
        for r in kpi_rows
    ]

    return summary


def approve_data(conn, domain: str, table: str, fiscal_year: int | None = None) -> int:
    """Set is_authoritative=1 for reviewed data. Returns count updated."""
    if table == "deal_customers":
        if fiscal_year:
            n = conn.execute(
                "UPDATE deal_customers SET is_authoritative = 1 WHERE domain = ? AND fiscal_year = ?",
                (domain, fiscal_year),
            ).rowcount
        else:
            n = conn.execute(
                "UPDATE deal_customers SET is_authoritative = 1 WHERE domain = ?",
                (domain,),
            ).rowcount
    elif table == "deal_commercial":
        if fiscal_year:
            n = conn.execute(
                "UPDATE deal_commercial SET is_authoritative = 1 WHERE domain = ? AND fiscal_year = ? AND category = 'customers'",
                (domain, fiscal_year),
            ).rowcount
        else:
            n = conn.execute(
                "UPDATE deal_commercial SET is_authoritative = 1 WHERE domain = ? AND category = 'customers'",
                (domain,),
            ).rowcount
    else:
        raise ValueError(f"Unsupported table for approval: {table}")

    conn.commit()
    return n


# ---- Main entry point ----------------------------------------------------


def ingest_commercial(
    conn,
    code_name: str,
    file_path: str | None = None,
    auto: bool = False,
    dry_run: bool = False,
    force: bool = False,
) -> dict:
    """Main entry: classify + confirm + load commercial files for a deal.

    Steps:
    1. Query deal_documents for commercial_raw files (or specific file)
    2. For each file:
       a. Check extraction_config -- if exists and not --force, use cached mapping
       b. If no config: classify via AI, print mapping, ask for confirm (unless --auto)
       c. Load data using confirmed mapping
       d. Derive KPIs
       e. Run cross-checks
    3. Return summary

    Returns: {files_processed, rows_written, kpis_derived, cross_checks}
    """
    deal = conn.execute(
        "SELECT domain, company_name FROM deals WHERE code_name = ?", (code_name,)
    ).fetchone()
    if not deal:
        raise ValueError(f"Deal not found: {code_name}")

    domain = deal["domain"] or code_name.lower()

    # Query commercial_raw files
    if file_path:
        docs = conn.execute(
            """SELECT id, file_path, file_name, doc_subtype, fiscal_year, extraction_config
               FROM deal_documents
               WHERE code_name = ? AND file_path = ?""",
            (code_name, file_path),
        ).fetchall()
    else:
        docs = conn.execute(
            """SELECT id, file_path, file_name, doc_subtype, fiscal_year, extraction_config
               FROM deal_documents
               WHERE code_name = ? AND doc_type = 'commercial_raw'
               AND (file_name LIKE '%.xlsx' OR file_name LIKE '%.xlsm')
               ORDER BY doc_subtype, fiscal_year""",
            (code_name,),
        ).fetchall()

    if not docs:
        return {
            "files_processed": 0,
            "rows_written": 0,
            "kpis_derived": 0,
            "cross_checks": [],
            "error": "No commercial_raw files found. Run ingest-docs first.",
        }

    result = {
        "files_processed": 0,
        "rows_written": 0,
        "kpis_derived": 0,
        "cross_checks": [],
        "errors": [],
        "skipped": [],
    }

    for doc in docs:
        doc_id = doc["id"]
        fname = doc["file_name"]
        safe_name = fname.encode("ascii", "replace").decode("ascii")
        path = Path(doc["file_path"])

        if not path.exists():
            result["skipped"].append(f"{safe_name} [file not found]")
            continue

        print(f"  Processing {safe_name} ({doc['doc_subtype']})...")

        # Step 2a: Check for cached mapping
        mapping = get_mapping(conn, doc_id)

        if mapping and not force:
            print(
                f"    Using cached mapping (target: {mapping.get('target_table', '?')})"
            )
        else:
            # Step 2b: Classify via AI
            try:
                print("    Classifying via Claude CLI...")
                mapping = classify_file(conn, doc_id)
            except Exception as e:
                result["errors"].append(f"{safe_name}: classification failed: {e}")
                continue

            # Print proposed mapping for review
            print("    Proposed mapping:")
            print(f"      file_type: {mapping.get('file_type')}")
            print(f"      target_table: {mapping.get('target_table')}")
            print(f"      fiscal_year: {mapping.get('fiscal_year')}")
            sheets = mapping.get("sheets", {})
            for sheet_name, config in sheets.items():
                print(f"      sheet: {sheet_name}")
                print(
                    f"        header_row: {config.get('header_row')}, data_start_row: {config.get('data_start_row')}, skip_footer: {config.get('skip_footer_rows', 0)}"
                )
                for cm in config.get("column_map", []):
                    print(
                        f"        {cm['col']} -> {cm['target']} ({cm.get('transform', 'text')})"
                    )

            if not auto:
                # Interactive: ask for confirmation
                answer = input("    Accept mapping? [y/n/q]: ").strip().lower()
                if answer == "q":
                    print("    Aborted.")
                    return result
                if answer != "y":
                    result["skipped"].append(f"{safe_name} [mapping rejected]")
                    continue

            # Save confirmed mapping
            confirm_mapping(conn, doc_id, mapping)
            print("    Mapping saved.")

        # Step 2c: Load data
        try:
            load_result = load_file(conn, doc_id, dry_run=dry_run)
            n = load_result["rows_written"]
            result["rows_written"] += n
            result["files_processed"] += 1
            print(
                f"    -> {n} rows {'(dry-run)' if dry_run else 'written'} to {load_result['table']}"
                f" (skipped {load_result['rows_skipped']})"
            )
        except NotImplementedError as e:
            result["skipped"].append(f"{safe_name} [{e}]")
            continue
        except Exception as e:
            result["errors"].append(f"{safe_name}: load failed: {e}")
            continue

    # Step 2d: Derive KPIs (after all files loaded)
    if result["files_processed"] > 0 and not dry_run:
        try:
            kpi_result = derive_kpis(conn, domain, dry_run=dry_run)
            result["kpis_derived"] = kpi_result["kpis_written"]
            print(
                f"\n  KPIs derived: {kpi_result['kpis_written']} for years {kpi_result['fiscal_years']}"
            )
        except Exception as e:
            result["errors"].append(f"KPI derivation failed: {e}")

    # Step 2e: Cross-checks (AP-4: always verify)
    if result["files_processed"] > 0:
        # Get all fiscal years with customer data
        year_rows = conn.execute(
            "SELECT DISTINCT fiscal_year FROM deal_customers WHERE domain = ? ORDER BY fiscal_year",
            (domain,),
        ).fetchall()
        for yr_row in year_rows:
            yr = yr_row["fiscal_year"]
            check = cross_check_revenue(conn, domain, yr)
            if check:
                result["cross_checks"].append(check)
                status_icon = {
                    "ok": "OK",
                    "warning": "WARN",
                    "mismatch": "MISMATCH",
                }.get(check["status"], check["status"])
                delta_str = (
                    f"{check['delta_pct']:.1f}%"
                    if check.get("delta_pct") is not None
                    else "n/a"
                )
                print(
                    f"  Cross-check {yr}: [{status_icon}] delta={delta_str}"
                    f" (customers={check['customer_total_k']:.0f}K,"
                    f" gesamtleistung={check.get('gesamtleistung_k', 'n/a')}K)"
                )

    return result
