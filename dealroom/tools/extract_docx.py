"""Extract text from one or more .docx files. Reusable deal tool.

Usage:
    python extract_docx.py <path1> [path2] ...
    python extract_docx.py <path1> --output extracted.txt
    python extract_docx.py <path1> --tables   (include table content)
"""

import sys
import os

try:
    from docx import Document
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document


def extract(path, include_tables=False):
    doc = Document(path)
    lines = []
    lines.append(f"{'=' * 80}")
    lines.append(f"=== {os.path.basename(path)} ===")
    lines.append(f"=== {path} ===")
    lines.append(f"{'=' * 80}")

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)

    if include_tables:
        for i, table in enumerate(doc.tables):
            lines.append(f"\n--- Table {i + 1} ---")
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                lines.append(" | ".join(cells))

    return "\n".join(lines)


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(1)

    include_tables = "--tables" in args
    args = [a for a in args if a != "--tables"]

    output_file = None
    if "--output" in args:
        idx = args.index("--output")
        output_file = args[idx + 1]
        args = args[:idx] + args[idx + 2 :]

    results = []
    for path in args:
        if not os.path.exists(path):
            results.append(f"ERROR: File not found: {path}")
            continue
        try:
            results.append(extract(path, include_tables))
        except Exception as e:
            results.append(f"ERROR extracting {path}: {e}")

    output = "\n\n".join(results)
    if output_file:
        with open(output_file, "w", encoding="utf-8") as f:
            f.write(output)
        print(f"Written to {output_file}")
    else:
        print(output)


if __name__ == "__main__":
    main()
