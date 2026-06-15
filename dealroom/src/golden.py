"""Shared golden corpus loader for DEALROOM generators.

Reads Roman-curated reference documents from config/golden/{doc_type}/
and returns extracted text for use as few-shot examples in prompts.
"""

from __future__ import annotations

import json
from pathlib import Path

from config import settings


def _read_manifest() -> dict:
    manifest_path = settings.GOLDEN_CORPUS_DIR / "manifest.json"
    if not manifest_path.exists():
        return {"files": [], "gaps": []}
    return json.loads(manifest_path.read_text(encoding="utf-8"))


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".md":
        return path.read_text(encoding="utf-8")

    if suffix == ".docx":
        try:
            from docx import Document

            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception:
            return ""

    if suffix == ".pdf":
        try:
            from pdfplumber import open as open_pdf

            lines = []
            with open_pdf(str(path)) as pdf:
                for page in pdf.pages[:20]:
                    text = page.extract_text()
                    if text:
                        lines.append(text)
            return "\n".join(lines)
        except Exception:
            return ""

    if suffix == ".pptx":
        try:
            from pptx import Presentation

            prs = Presentation(str(path))
            lines = []
            for i, slide in enumerate(prs.slides, 1):
                slide_texts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            t = para.text.strip()
                            if t:
                                slide_texts.append(t)
                if slide_texts:
                    lines.append(f"--- Slide {i} ---")
                    lines.extend(slide_texts)
            return "\n".join(lines)
        except Exception:
            return ""

    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(path), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets[:3]:
                lines.append(f"--- Sheet: {ws.title} ---")
                for row in ws.iter_rows(max_row=50, values_only=True):
                    cells = [str(c) if c is not None else "" for c in row]
                    if any(cells):
                        lines.append("\t".join(cells))
            wb.close()
            return "\n".join(lines)
        except Exception:
            return ""

    return ""


def load_golden_corpus(
    doc_type: str,
    exclude_deal: str = "",
    max_chars: int = 6000,
) -> list[dict]:
    """Load golden files for a doc_type.

    Returns list of {path, text, metadata} dicts.
    Respects manifest.json for metadata.
    Excludes files from exclude_deal (avoid self-reference).
    Truncates total text to max_chars.
    """
    manifest = _read_manifest()
    results = []
    total_chars = 0

    for entry in manifest.get("files", []):
        if entry.get("exclude_from_corpus"):
            continue
        if entry.get("doc_type") != doc_type:
            continue
        if (
            exclude_deal
            and entry.get("source_deal", "").lower() == exclude_deal.lower()
        ):
            continue

        file_path = settings.GOLDEN_CORPUS_DIR / entry["path"]
        if not file_path.exists():
            continue

        text = _extract_text(file_path)
        if not text:
            continue

        remaining = max_chars - total_chars
        if remaining <= 0:
            break
        if len(text) > remaining:
            text = text[:remaining]

        total_chars += len(text)
        results.append(
            {
                "path": entry["path"],
                "text": text,
                "metadata": entry,
            }
        )

    return results


def load_golden_text(
    doc_type: str,
    exclude_deal: str = "",
    max_chars: int = 6000,
) -> str:
    """Convenience: load golden corpus as a single concatenated string."""
    entries = load_golden_corpus(doc_type, exclude_deal, max_chars)
    parts = []
    for e in entries:
        label = e["metadata"].get("source_deal", "unknown")
        parts.append(f"--- {label} / {Path(e['path']).name} ---\n{e['text']}")
    return "\n\n".join(parts)


def check_golden_health() -> dict:
    """Validate golden corpus. Returns {total, by_type, missing, gaps}."""
    manifest = _read_manifest()
    doc_types = [
        "rfi",
        "offer",
        "onepager",
        "email",
        "nda",
        "model",
        "databook",
        "loi",
        "dd",
    ]

    by_type: dict[str, list[str]] = {dt: [] for dt in doc_types}
    missing: list[str] = []

    for entry in manifest.get("files", []):
        dt = entry.get("doc_type", "unknown")
        file_path = settings.GOLDEN_CORPUS_DIR / entry["path"]
        if file_path.exists():
            by_type.setdefault(dt, []).append(entry["path"])
        else:
            missing.append(entry["path"])

    gaps = [dt for dt in doc_types if not by_type.get(dt)]
    manifest_gaps = manifest.get("gaps", [])

    return {
        "total": sum(len(v) for v in by_type.values()),
        "by_type": {k: len(v) for k, v in by_type.items()},
        "missing": missing,
        "gaps": gaps,
        "manifest_gaps": manifest_gaps,
    }
