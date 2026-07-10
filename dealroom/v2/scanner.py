"""DEALROOM v2 — local scanner CLI (spec M2): parses OneDrive locally, pushes
payloads to the dealroom API (spec §3: ingestion stays local, storage remote).

Commands (run from dealroom/):
    python -m v2.scanner artifacts --deal Fox [--dry-run]
    python -m v2.scanner dataroom  --deal Fox --path <dataroom-root> [--dry-run]
    python -m v2.scanner rfi       --deal Fox [--file <fragenliste.xlsx>] [--dry-run]
    python -m v2.scanner all       --deal Fox

API target: --api (default http://127.0.0.1:8082, the local sandbox server).
At Fly cutover the same commands push to prod with basic-auth credentials.

READ-ONLY GUARANTEE: this module only ever stat()s and reads files. Data-room
folders (Sorgfaltsprüfung shares) are never written — hard rule.

Per-deal path overrides live in data/scanner_paths.json:
    {"dealrooms": {"Fox": "C:/path/to/dataroom-root"},
     "targets_root": "optional override"}
"""

import argparse
import json
import re
import sys
import urllib.error
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from v2 import db, naming  # noqa: E402

DEFAULT_TARGETS_ROOT = Path(
    r"C:\Users\X1\Documents\OneDrive - Kamu Kapital\Dokumente - Kamu Kapital"
    r"\3_Deals\3_Targets"
)
SECTION_RE = re.compile(r"^(0[0-8])[_ ]")

CONFIG_FILE = Path(__file__).resolve().parent.parent / "data" / "scanner_paths.json"


def load_config() -> dict:
    if CONFIG_FILE.exists():
        return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    return {}


def iso_mtime(p: Path) -> str:
    return (
        datetime.fromtimestamp(p.stat().st_mtime, tz=timezone.utc)
        .replace(microsecond=0)
        .isoformat()
        .replace("+00:00", "Z")
    )


def api_post(api: str, path: str, payload: dict, dry_run: bool) -> dict:
    if dry_run:
        print(f"[dry-run] POST {path}")
        print(json.dumps(payload, ensure_ascii=False, indent=1)[:2000])
        return {"ok": True, "dry_run": True}
    req = urllib.request.Request(
        api.rstrip("/") + path,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            return json.loads(resp.read())
    except urllib.error.HTTPError as e:
        sys.exit(f"API error {e.code} on {path}: {e.read().decode()[:400]}")
    except urllib.error.URLError as e:
        sys.exit(
            f"API unreachable ({e.reason}) — is the v2 server running? "
            "Start: python -m v2.server --port 8082"
        )


def deal_folder(code: str, cfg: dict) -> Path:
    root = Path(cfg.get("targets_root") or DEFAULT_TARGETS_ROOT)
    conn = db.open_readonly(db.default_db_path())
    row = conn.execute(
        "SELECT folder_path FROM deals WHERE lower(code_name)=lower(?)", (code,)
    ).fetchone()
    conn.close()
    if not row or not row["folder_path"]:
        sys.exit(f"no folder_path for deal {code!r} in v2 DB")
    folder = root / row["folder_path"]
    if not folder.exists():
        sys.exit(f"deal folder not found: {folder}")
    return folder


# ------------------------------------------------------------- artifacts ----


def collect_artifacts(folder: Path) -> list[dict]:
    out = []
    for p in sorted(folder.rglob("*")):
        if not p.is_file() or naming.skip_file(p.name):
            continue
        rel = p.relative_to(folder).as_posix()
        if any(seg.lower() in naming.SKIP_DIRS for seg in rel.split("/")[:-1]):
            continue
        out.append(
            {
                "file_path": str(p),
                "file_name": p.name,
                "artifact_type": naming.classify(p.name, rel),
                "version": naming.parse_version(p.name),
                "file_date": naming.parse_file_date(p.name),
                "file_mtime": iso_mtime(p),
                "file_size_kb": max(1, p.stat().st_size // 1024),
            }
        )
    return out


def cmd_artifacts(args, cfg):
    folder = deal_folder(args.deal, cfg)
    artifacts = collect_artifacts(folder)
    by_type = {}
    for a in artifacts:
        by_type[a["artifact_type"]] = by_type.get(a["artifact_type"], 0) + 1
    print(f"{args.deal}: {len(artifacts)} files in {folder}")
    print("  " + ", ".join(f"{k}={v}" for k, v in sorted(by_type.items())))
    res = api_post(
        args.api,
        "/api/push/artifacts",
        {"code_name": args.deal, "artifacts": artifacts},
        args.dry_run,
    )
    print(f"  pushed: {res}")


# -------------------------------------------------------------- dataroom ----


def cmd_dataroom(args, cfg):
    path = args.path or (cfg.get("dealrooms") or {}).get(args.deal)
    if not path:
        sys.exit(
            f"no data-room path for {args.deal!r} — pass --path or add it to "
            f"{CONFIG_FILE}"
        )
    root = Path(path)
    if not root.exists():
        sys.exit(f"data-room root not found: {root}")
    sections = []
    for d in sorted(root.iterdir()):
        m = SECTION_RE.match(d.name) if d.is_dir() else None
        if not m:
            continue
        files = []
        for p in sorted(d.rglob("*")):
            if p.is_file() and not naming.skip_file(p.name):
                files.append(
                    {"name": p.relative_to(d).as_posix(), "mtime": iso_mtime(p)}
                )
        newest = max(files, key=lambda f: f["mtime"], default=None)
        sections.append(
            {
                "section": m.group(1),
                "section_name": d.name[len(m.group(0)) :].strip("_ "),
                "files": files,
                "newest_file_date": newest["mtime"][:10] if newest else None,
                "newest_file_name": newest["name"] if newest else None,
            }
        )
    if not sections:
        sys.exit(f"no 00–08 section folders under {root}")
    print(
        f"{args.deal}: {len(sections)} sections, "
        f"{sum(len(s['files']) for s in sections)} files"
    )
    res = api_post(
        args.api,
        "/api/push/dataroom-scan",
        {"code_name": args.deal, "sections": sections},
        args.dry_run,
    )
    print(
        f"  pushed: {res if args.dry_run else {k: res[k] for k in ('ok', 'sections')}}"
    )


# ------------------------------------------------------------------- rfi ----

RFI_NAME_RE = re.compile(r"fragenliste|datenanfrage", re.IGNORECASE)
PRIO_MAP = {
    "hoch": "high",
    "mittel": "medium",
    "niedrig": "low",
    "high": "high",
    "medium": "medium",
    "low": "low",
}


def find_rfi_file(folder: Path) -> Path | None:
    candidates = (
        [
            p
            for p in (folder / "5_DD").glob("*.xlsx")
            if RFI_NAME_RE.search(p.name) and not naming.skip_file(p.name)
        ]
        if (folder / "5_DD").exists()
        else []
    )
    if not candidates:
        return None
    return max(
        candidates,
        key=lambda p: (
            naming.parse_file_date(p.name) or "",
            "_vs" in p.name.lower(),  # versandt beats print/draft on same date
            p.stat().st_mtime,
        ),
    )


def parse_fragenliste(path: Path) -> list[dict]:
    """Parse the Repuro Fragenliste convention (header row: Nr. | [Thema] |
    Frage | Priorität | [Quelle] [Status] [Antwort ...] [Kommentar])."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = next(
            (w for w in wb.worksheets if "fragen" in w.title.lower()),
            wb.worksheets[0],
        )
        header_idx, header = None, {}
        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
        for i, row in enumerate(rows[:12]):
            cells = [str(c).strip() if c else "" for c in row]
            if "Frage" in cells:  # Priorität column is optional (Mantis vS)
                header_idx = i
                for j, c in enumerate(cells):
                    header[c.lower().split(" ")[0]] = j
                break
        if header_idx is None:
            sys.exit(f"no Fragenliste header row found in {path.name}")

        col_frage = header.get("frage")
        col_prio = next((v for k, v in header.items() if k.startswith("priorit")), None)
        col_thema = header.get("thema", header.get("quelle"))
        col_status = header.get("status")
        col_antwort = next(
            (v for k, v in header.items() if k.startswith("antwort")), None
        )
        sent_date = naming.parse_file_date(path.name)

        questions = []
        for row in rows[header_idx + 1 :]:
            frage = row[col_frage] if col_frage is not None else None
            if not frage or not str(frage).strip():
                continue
            answer = (
                str(row[col_antwort]).strip()
                if col_antwort is not None and row[col_antwort]
                else None
            )
            raw_status = (
                str(row[col_status]).strip().lower()
                if col_status is not None and row[col_status]
                else ""
            )
            if answer:
                status = "answered"
            elif "beantwortet" in raw_status or "erledigt" in raw_status:
                status = "answered"
            else:
                status = "sent"
            prio_raw = (
                str(row[col_prio]).strip().lower()
                if col_prio is not None and row[col_prio]
                else ""
            )
            questions.append(
                {
                    "question": str(frage).strip(),
                    "category": (
                        str(row[col_thema]).strip()
                        if col_thema is not None and row[col_thema]
                        else "RFI"
                    ),
                    "importance": PRIO_MAP.get(prio_raw, "medium"),
                    "status": status,
                    "answer": answer,
                    "sent_at": sent_date,
                    "source": path.name,
                }
            )
        return questions
    finally:
        wb.close()


def cmd_rfi(args, cfg):
    folder = deal_folder(args.deal, cfg)
    path = Path(args.file) if args.file else find_rfi_file(folder)
    if not path or not path.exists():
        sys.exit(f"no Fragenliste xlsx found for {args.deal} (5_DD)")
    questions = parse_fragenliste(path)
    n_high = sum(1 for q in questions if q["importance"] == "high")
    n_open = sum(1 for q in questions if q["status"] == "sent")
    print(
        f"{args.deal}: {len(questions)} Fragen aus {path.name} "
        f"({n_high} HIGH, {n_open} offen)"
    )
    res = api_post(
        args.api,
        "/api/push/rfi",
        {"code_name": args.deal, "questions": questions},
        args.dry_run,
    )
    print(f"  pushed: {res}")


# ------------------------------------------------------------------ main ----


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["artifacts", "dataroom", "rfi", "all"])
    parser.add_argument("--deal", required=True)
    parser.add_argument("--api", default="http://127.0.0.1:8082")
    parser.add_argument("--path", help="data-room root (dataroom command)")
    parser.add_argument("--file", help="explicit Fragenliste xlsx (rfi command)")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    cfg = load_config()

    if args.command in ("artifacts", "all"):
        cmd_artifacts(args, cfg)
    if args.command in ("rfi", "all"):
        cmd_rfi(args, cfg)
    if args.command == "dataroom" or (
        args.command == "all" and (cfg.get("dealrooms") or {}).get(args.deal)
    ):
        cmd_dataroom(args, cfg)


if __name__ == "__main__":
    main()
