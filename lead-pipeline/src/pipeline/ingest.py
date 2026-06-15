"""Data ingestion — normalize all sources to CompanyRecord, write to pipeline.db."""

from __future__ import annotations

import csv
import logging
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Optional, Protocol, runtime_checkable

import openpyxl

from src.config import settings
from src.config.profile import IndustryProfile
from src.pipeline import db as pipeline_db
from src.pipeline.models import CompanyRecord

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Domain normalization
# ---------------------------------------------------------------------------

_DOMAIN_PREFIX_RE = re.compile(r"^(https?://)?(www\.)?", re.IGNORECASE)


def normalize_domain(raw: Optional[str]) -> Optional[str]:
    """Strip protocol, www, path, trailing slash. Lowercase. Return None if empty."""
    if not raw or not str(raw).strip():
        return None
    s = str(raw).strip().lower()
    s = _DOMAIN_PREFIX_RE.sub("", s)
    s = s.split("/")[0].strip()
    return s or None


# ---------------------------------------------------------------------------
# Source adapter protocol
# ---------------------------------------------------------------------------


@runtime_checkable
class SourceAdapter(Protocol):
    @property
    def name(self) -> str: ...
    def read(self) -> list[CompanyRecord]: ...


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_CORPORATE_KW = frozenset(kw.lower() for kw in settings.CORPORATE_KEYWORDS)


def _is_corporate_name(name: Optional[str]) -> bool:
    if not name:
        return False
    low = name.lower()
    return any(kw in low for kw in _CORPORATE_KW)


def _excel_serial_to_age(serial: object) -> Optional[int]:
    """Convert Excel date serial number to approximate current age in years."""
    if serial is None:
        return None
    try:
        n = float(str(serial))  # type: ignore[arg-type]
    except (ValueError, TypeError):
        return None
    if n <= 0:
        return None
    birth_date = datetime(1899, 12, 30) + timedelta(days=n)
    age = (datetime.now() - birth_date).days // 365
    return age if 0 < age < 120 else None


def _safe_float(val: object) -> Optional[float]:
    if val is None or str(val).strip() in ("", "n.v.", "N/A", "nan"):
        return None
    try:
        return float(str(val).replace(",", ".").strip())
    except (ValueError, TypeError):
        return None


def _safe_int(val: object) -> Optional[int]:
    f = _safe_float(val)
    return int(f) if f is not None else None


# ---------------------------------------------------------------------------
# Adapter: WLW CSV
# ---------------------------------------------------------------------------


class WLWCsvAdapter:
    """Reads data/staging/00_wlw_raw.csv — already in CompanyRecord column format."""

    def __init__(self, path: Path) -> None:
        self._path = path

    @property
    def name(self) -> str:
        return "WLW CSV"

    def read(self) -> list[CompanyRecord]:
        if not self._path.exists():
            logger.warning("WLW CSV not found: %s — skipping", self._path)
            return []
        records: list[CompanyRecord] = []
        with open(self._path, encoding="utf-8", newline="") as f:
            for row in csv.DictReader(f):
                domain = normalize_domain(row.get("domain"))
                if not domain:
                    continue
                row["domain"] = domain
                try:
                    records.append(CompanyRecord.from_dict(row))
                except Exception as exc:
                    logger.debug("Skipping WLW row (parse error): %s", exc)
        logger.info("WLW CSV: %d records loaded", len(records))
        return records


# ---------------------------------------------------------------------------
# Adapter: ORBIS sheet (reads from pre-opened openpyxl worksheet)
# ---------------------------------------------------------------------------


class OrbisSheetAdapter:
    """Reads ORBIS_search worksheet → CompanyRecord with ownership pre-population."""

    def __init__(self, ws: object, profile_id: str) -> None:
        self._ws = ws
        self._profile_id = profile_id

    @property
    def name(self) -> str:
        return "ORBIS Excel"

    def read(self) -> list[CompanyRecord]:
        records: list[CompanyRecord] = []
        skipped = 0

        for row in self._ws.iter_rows(min_row=2, values_only=True):
            domain = normalize_domain(row[settings.ORBIS_COL_DOMAIN])
            if not domain:
                skipped += 1
                continue

            full_name = str(row[settings.ORBIS_COL_NAME] or "").strip() or domain

            plz_raw = str(row[settings.ORBIS_COL_PLZ] or "").strip()
            city_raw = str(row[settings.ORBIS_COL_CITY] or "").strip()
            plz_ort = f"{plz_raw} {city_raw}".strip() if (plz_raw or city_raw) else None

            vorname = str(row[settings.ORBIS_COL_GF_VORNAME] or "").strip()
            nachname = str(row[settings.ORBIS_COL_GF_NACHNAME] or "").strip()
            if vorname or nachname:
                owner_name: Optional[str] = f"{vorname} {nachname}".strip() or None
            else:
                dm_last = str(row[settings.ORBIS_COL_DM_NACHNAME] or "").strip()
                owner_name = dm_last or None

            anrede_raw = str(row[settings.ORBIS_COL_ANREDE] or "").strip()
            anrede: Optional[str] = anrede_raw if anrede_raw else None

            ges_name_raw = str(
                row[settings.ORBIS_COL_GESELLSCHAFTER_NAME] or ""
            ).strip()
            csh_name_raw = str(row[settings.ORBIS_COL_CSH_NAME] or "").strip()
            ges_pct_raw = row[settings.ORBIS_COL_GESELLSCHAFTER_PCT]
            csh_total_raw = row[settings.ORBIS_COL_CSH_TOTAL_PCT]

            if ges_name_raw:
                gesellschafter_name: Optional[str] = ges_name_raw
                gesellschafter_share_pct = _safe_float(ges_pct_raw)
            elif csh_name_raw:
                gesellschafter_name = csh_name_raw
                gesellschafter_share_pct = _safe_float(csh_total_raw)
            else:
                gesellschafter_name = None
                gesellschafter_share_pct = None

            gesellschafter_age = _excel_serial_to_age(
                row[settings.ORBIS_COL_DM_BIRTHDAY]
            )

            csh_total_pct = _safe_float(csh_total_raw)
            is_subsidiary: Optional[bool] = None
            if csh_name_raw and csh_total_pct is not None and csh_total_pct >= 50.0:
                if _is_corporate_name(csh_name_raw):
                    is_subsidiary = True

            records.append(
                CompanyRecord(
                    domain=domain,
                    full_name=full_name,
                    profile_id=self._profile_id,
                    source="ORBIS",
                    street=str(row[settings.ORBIS_COL_STREET] or "").strip() or None,
                    plz_ort=plz_ort,
                    city=city_raw or None,
                    ma_count=_safe_int(row[settings.ORBIS_COL_MA]),
                    revenue_tsd_eur=_safe_float(row[settings.ORBIS_COL_REVENUE]),
                    owner_name=owner_name,
                    anrede=anrede,
                    gesellschafter_name=gesellschafter_name,
                    gesellschafter_share_pct=gesellschafter_share_pct,
                    gesellschafter_age=gesellschafter_age,
                    is_subsidiary=is_subsidiary,
                )
            )

        logger.info(
            "ORBIS: %d records loaded, %d skipped (no domain)", len(records), skipped
        )
        return records


# ---------------------------------------------------------------------------
# LLMPrep enrichment (reads from pre-opened worksheet)
# ---------------------------------------------------------------------------


def _read_llm_prep_enrichment(ws: object) -> dict[str, dict]:
    """Returns {domain: {hrb_number, owner_name, ma_count}} for patching records."""
    enrichment: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=settings.LLM_PREP_HEADER_ROW + 1, values_only=True):
        domain = normalize_domain(row[settings.LLM_PREP_COL_DOMAIN])
        if not domain:
            continue
        hrb_raw = str(row[settings.LLM_PREP_COL_HRB] or "").strip()
        gf_raw = str(row[settings.LLM_PREP_COL_GF] or "").strip()
        owner_name = gf_raw.split(",")[0].strip() if gf_raw else None
        enrichment[domain] = {
            "hrb_number": hrb_raw or None,
            "owner_name": owner_name,
            "ma_count": _safe_int(row[settings.LLM_PREP_COL_MA]),
        }
    logger.info("LLMPrep: %d enrichment entries loaded", len(enrichment))
    return enrichment


# ---------------------------------------------------------------------------
# Serienbriefe dedup (reads from pre-opened worksheet)
# ---------------------------------------------------------------------------

_SB_LEGAL_FORMS = re.compile(
    r"\b(gmbh\s*&\s*co\.?\s*kg|gmbh\s*&\s*co|gmbh|ag|se|kg|gbr|ohg|ug|inc|ltd|plc|corp|group|holding|gesellschaft|mbh)\b",
    re.IGNORECASE,
)


def _normalize_company_name(name: str) -> str:
    """Normalize company name for fuzzy dedup (strip legal forms, punctuation)."""
    if not name:
        return ""
    s = str(name).lower()
    s = _SB_LEGAL_FORMS.sub("", s)
    s = re.sub(r"[^a-z0-9äöüß]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _read_serienbriefe_dedup(ws: object) -> tuple[frozenset[str], dict[str, str]]:
    """Load already-approached domains and company names from Serienbriefe sheet.

    Returns:
        domains: frozenset of normalized domains (col 2)
        name_index: {normalized_company_name: domain} for name-based fallback matching
                    Names come from col 11 (Name Briefkopf). Min key length 4 chars.
    """
    domains: set[str] = set()
    name_index: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = normalize_domain(row[settings.SERIENBRIEFE_COL_DOMAIN])
        if d:
            domains.add(d)
        name_raw = row[11] if len(row) > 11 else None
        name_key = _normalize_company_name(str(name_raw or ""))
        if name_key and len(name_key) >= 4 and d:
            name_index[name_key] = d
    logger.info(
        "Serienbriefe: %d domains + %d name-index entries loaded",
        len(domains),
        len(name_index),
    )
    return frozenset(domains), name_index


# ---------------------------------------------------------------------------
# Main ingest function
# ---------------------------------------------------------------------------


def ingest(
    profile: IndustryProfile,
    dry_run: bool = False,
    verbose: bool = False,
) -> int:
    """Load all sources, normalize, write to pipeline.db. Returns record count written."""

    ingested_at = datetime.now(timezone.utc).isoformat()

    # 1. WLW CSV (M2 scraper output — no Excel dependency)
    wlw_records = WLWCsvAdapter(settings.STAGING_WLW_RAW).read()

    # 2. Open source Excel ONCE — read all three sheets in one pass
    orbis_records: list[CompanyRecord] = []
    llm_enrichment: dict[str, dict] = {}
    approached_domains: frozenset[str] = frozenset()
    approached_names: dict[str, str] = {}

    if settings.SOURCE_EXCEL.exists():
        logger.info("Opening source Excel (once): %s", settings.SOURCE_EXCEL.name)
        wb = openpyxl.load_workbook(
            settings.SOURCE_EXCEL, read_only=True, data_only=True
        )
        try:
            orbis_records = OrbisSheetAdapter(
                wb[settings.ORBIS_SHEET], profile.id
            ).read()
            llm_enrichment = _read_llm_prep_enrichment(wb[settings.LLM_PREP_SHEET])
            approached_domains, approached_names = _read_serienbriefe_dedup(
                wb[settings.SERIENBRIEFE_SHEET]
            )
        finally:
            wb.close()
        logger.info("Excel closed.")
    else:
        logger.warning(
            "Source Excel not found: %s — skipping ORBIS/LLMPrep/Serienbriefe",
            settings.SOURCE_EXCEL,
        )

    total_sources = len(wlw_records) + len(orbis_records)
    logger.info("Total from all sources (before dedup): %d", total_sources)

    if dry_run:
        logger.info(
            "DRY RUN — skipping DB write. Would process %d records.", total_sources
        )
        return total_sources

    # 3. Write to pipeline.db
    # Insert order: ORBIS first (higher priority for structured fields).
    # The COALESCE upsert keeps the first non-null value, so ORBIS wins.
    pipeline_db.ensure_schema(settings.PIPELINE_DB_PATH)

    with pipeline_db.get_connection(settings.PIPELINE_DB_PATH) as conn:
        # ORBIS records — apply LLMPrep enrichment before insert
        def _is_approached(rec: "CompanyRecord") -> bool:
            """Domain match OR normalized name match against Serienbriefe."""
            if rec.domain in approached_domains:
                return True
            name_key = _normalize_company_name(rec.full_name)
            return bool(
                name_key and len(name_key) >= 4 and name_key in approached_names
            )

        for rec in orbis_records:
            patch = llm_enrichment.get(rec.domain)
            if patch:
                if rec.hrb_number is None and patch.get("hrb_number"):
                    object.__setattr__(rec, "hrb_number", patch["hrb_number"])
                if rec.owner_name is None and patch.get("owner_name"):
                    object.__setattr__(rec, "owner_name", patch["owner_name"])
                if rec.ma_count is None and patch.get("ma_count"):
                    object.__setattr__(rec, "ma_count", patch["ma_count"])
            if _is_approached(rec):
                object.__setattr__(rec, "already_approached", True)
            pipeline_db.upsert_company(conn, rec, ingested_at)

        # WLW records — fills scraped_text and other nulls into ORBIS records
        for rec in wlw_records:
            if _is_approached(rec):
                object.__setattr__(rec, "already_approached", True)
            pipeline_db.upsert_company(conn, rec, ingested_at)

        # Final sweep: mark any remaining approached domains (domain-only, belt-and-suspenders)
        approached_count = pipeline_db.mark_approached(conn, approached_domains)

    total = pipeline_db.get_total_count(settings.PIPELINE_DB_PATH)
    logger.info(
        "pipeline.db: %d unique records after ingest (%d already-approached)",
        total,
        approached_count,
    )
    return total


# ---------------------------------------------------------------------------
# Serienbriefe full ingest (M14)
# ---------------------------------------------------------------------------

# Send date (ISO) → Briefaktion label, ordered chronologically
_BRIEFAKTION_DATE_MAP: dict[str, str] = {
    "2025-06-10": "BA1",
    "2025-06-17": "BA2",
    "2025-06-23": "BA3",
    "2025-07-09": "BA4",
    "2025-08-06": "BA5",
    "2025-10-15": "BA6",
    "2025-11-11": "BA7",
}

# Serienbriefe "Status / Follow-Up Status" → outreach_status
# Granular values preserve the full tracking chain from the sheet.
# Positive = contact, meeting, financials, offer, deal
# No Response = sent, followup1, followup2
# Declined = declined  |  Neutral = hold
_FU_STATUS_MAP: dict[Optional[str], str] = {
    None: "sent",
    "": "sent",
    "Follow-Up Sent": "followup1",
    "Follow-up Sent": "followup1",
    "Follow-Up 2 Sent": "followup2",
    "Contact Made": "contact",
    "Initial meeting done": "meeting",
    "Financials Received": "financials",
    "Indicative Offer Sent": "offer",
    "Converted to Deal": "deal",
    "Hold": "hold",
    "Lost": "declined",
}

# Column indices (0-based) in the Serienbriefe sheet
_C_DATE = 1  # Datum Sent
_C_DOMAIN = 2  # Domain Name Clean
_C_KLASS = 4  # Category
_C_MA = 5  # MA
_C_NAME_BK = 11  # Name Briefkopf (company display name)
_C_STREET = 18  # Street Address
_C_PLZ = 19  # PLZ + Stadt
_C_CITY = 20  # Stadt
_C_ANREDE = 25  # Anrede
_C_SALUTATION = 26  # Salutation
_C_GS_NAME = 30  # Name (1) — full Gesellschafter name
_C_EMAIL = 31  # Email
_C_PHONE = 32  # Tel
_C_AGE = 33  # Age
_C_HRB = 41  # HR-Nummer
_C_RECHTSFORM = 47  # Rechtsform
_C_K1 = 22  # Kompliment 1
_C_K2 = 23  # Kompliment 2
_C_FU_STATUS = 44  # Status / Follow-Up Status
_C_COMMENT = 45  # Comment


def _parse_sb_date(raw: object) -> tuple[Optional[str], Optional[str]]:
    """Return (iso_date, briefaktion) from a Datum Sent cell. Returns (None, None) on failure."""
    if raw is None:
        return None, None
    if isinstance(raw, datetime):
        iso = raw.date().isoformat()
    else:
        s = str(raw).strip()
        iso = None
        for fmt in ("%d/%m/%y", "%d/%m/%Y", "%Y-%m-%d"):
            try:
                iso = datetime.strptime(s, fmt).date().isoformat()
                break
            except ValueError:
                continue
        if iso is None:
            return None, None
    return iso, _BRIEFAKTION_DATE_MAP.get(iso)


def ingest_serienbriefe(dry_run: bool = False) -> dict:
    """Ingest the Serienbriefe sheet into pipeline.db as full approached records.

    New domains → INSERT with source='serienbriefe', pipeline_stage='approached'.
    Existing domains → UPDATE: briefaktion always set; outreach/enrichment fields
    filled via COALESCE (won't overwrite existing pipeline data).

    Returns stats: total, inserted_new, updated_existing,
                   skipped_no_domain, skipped_no_date, skipped_unknown_ba.
    """
    import hashlib as _hashlib
    from src.pipeline.category_defaults import (
        TARGET_CATEGORIES,
        NON_TARGET_CATEGORIES,
    )
    from src.pipeline.db import (
        ensure_schema,
        get_connection,
        upsert_serienbriefe_record,
    )

    _ALL_CATEGORIES = TARGET_CATEGORIES | NON_TARGET_CATEGORIES
    # Legacy codes still accepted for backwards compat
    _LEGACY_KLASS = {"A", "B", "C", "D", "E", "S"}
    _VALID_KLASS = _ALL_CATEGORIES | _LEGACY_KLASS

    db_path = settings.PIPELINE_DB_PATH

    if not dry_run:
        ensure_schema(db_path)

    wb = openpyxl.load_workbook(settings.SOURCE_EXCEL, read_only=True, data_only=True)
    ws = wb[settings.SERIENBRIEFE_SHEET]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    stats: dict[str, int] = {
        "total": len(rows),
        "inserted_new": 0,
        "updated_existing": 0,
        "skipped_no_domain": 0,
        "skipped_no_date": 0,
        "skipped_unknown_ba": 0,
    }

    ingested_at = datetime.now(timezone.utc).isoformat()
    to_process: list[dict] = []

    for row in rows:
        domain = normalize_domain(row[_C_DOMAIN])
        if not domain:
            stats["skipped_no_domain"] += 1
            continue

        date_iso, briefaktion = _parse_sb_date(row[_C_DATE])
        if not date_iso:
            stats["skipped_no_date"] += 1
            continue
        if not briefaktion:
            logger.warning("Unknown BA for date %s (domain: %s)", date_iso, domain)
            stats["skipped_unknown_ba"] += 1
            continue

        fu_raw = row[_C_FU_STATUS]
        outreach_status = _FU_STATUS_MAP.get(fu_raw, "sent")

        age_raw = row[_C_AGE]
        age = (
            int(age_raw)
            if isinstance(age_raw, (int, float)) and not isinstance(age_raw, bool)
            else None
        )

        phone_raw = row[_C_PHONE]
        phone = str(phone_raw).strip() if phone_raw and phone_raw != 0 else None

        email_raw = row[_C_EMAIL]
        email = str(email_raw).strip() if email_raw and str(email_raw).strip() else None

        ma_raw = row[_C_MA]
        ma = (
            int(ma_raw)
            if isinstance(ma_raw, (int, float)) and not isinstance(ma_raw, bool)
            else None
        )

        def _s(val: object) -> Optional[str]:
            return str(val).strip() if val else None

        # Accept legacy codes (A/B/C/E/S) and new BA9 category codes (DEA, INT, etc.)
        klass_raw = _s(row[_C_KLASS])
        klass_value = klass_raw if klass_raw in _VALID_KLASS else None

        anrede_value = _s(row[_C_ANREDE])

        rec = {
            "id": _hashlib.md5(domain.encode()).hexdigest(),
            "domain": domain,
            "full_name": _s(row[_C_NAME_BK]) or domain,
            "profile_id": "medtech_germany",
            "source": "serienbriefe",
            "pipeline_stage": "approached",
            "briefaktion": briefaktion,
            "outreach_status": outreach_status,
            "outreach_sent_at": date_iso,
            "outreach_comment": _s(row[_C_COMMENT]),
            "klass": klass_value,
            "owner_name": _s(row[_C_GS_NAME]),
            "gesellschafter_name": _s(row[_C_GS_NAME]),
            "gf_email": email,
            "gf_phone": phone,
            "gesellschafter_age": age,
            "anrede": anrede_value,
            "salutation": _s(row[_C_SALUTATION]),
            "hrb_number": _s(row[_C_HRB]),
            "rechtsform": _s(row[_C_RECHTSFORM]),
            "ma_count": ma,
            "city": _s(row[_C_CITY]),
            "street": _s(row[_C_STREET]),
            "plz_ort": _s(row[_C_PLZ]),
            "compliment_draft": _s(row[_C_K1]),
            "compliment_2": _s(row[_C_K2]),
        }

        # BA9+: apply prio + category defaults + gesellschafter_field
        if klass_value and klass_value in _ALL_CATEGORIES:
            from src.pipeline.category_defaults import (
                apply_category_defaults,
                get_default_prio,
                get_gesellschafter_field,
            )

            rec["prio"] = get_default_prio(klass_value)
            apply_category_defaults(rec, klass_value)
            if anrede_value:
                rec["gesellschafter_field"] = get_gesellschafter_field(anrede_value)
        elif klass_value and klass_value in _LEGACY_KLASS:
            if klass_value in ("A", "B"):
                rec["prio"] = "Prio 1"
            elif klass_value == "S":
                rec["prio"] = "Excluded"
            elif klass_value in ("C", "D", "E"):
                rec["prio"] = "Prio 2 (other)"

        to_process.append(rec)

    logger.info(
        "Serienbriefe: %d valid records to process (%d skipped)",
        len(to_process),
        stats["skipped_no_domain"]
        + stats["skipped_no_date"]
        + stats["skipped_unknown_ba"],
    )

    if dry_run:
        logger.info("DRY RUN — no writes performed")
        return stats

    with get_connection(db_path) as conn:
        for rec in to_process:
            result = upsert_serienbriefe_record(conn, rec, ingested_at)
            if result == "inserted":
                stats["inserted_new"] += 1
            else:
                stats["updated_existing"] += 1

            # BA9+: write category default fields that upsert_serienbriefe_record
            # does not handle (prio, leistung, mehrwerte, gruppe, gesellschafter_field)
            if rec.get("prio") or rec.get("leistung_text"):
                conn.execute(
                    """UPDATE company_records
                       SET prio = COALESCE(:prio, prio),
                           leistung_text = COALESCE(:leistung_text, leistung_text),
                           leistung_absatz_2 = COALESCE(:leistung_absatz_2, leistung_absatz_2),
                           mehrwerte = COALESCE(:mehrwerte, mehrwerte),
                           gruppe_1 = COALESCE(:gruppe_1, gruppe_1),
                           gruppe_2 = COALESCE(:gruppe_2, gruppe_2),
                           gesellschafter_field = COALESCE(:gesellschafter_field, gesellschafter_field)
                     WHERE domain = :domain""",
                    {
                        "domain": rec["domain"],
                        "prio": rec.get("prio"),
                        "leistung_text": rec.get("leistung_text"),
                        "leistung_absatz_2": rec.get("leistung_absatz_2"),
                        "mehrwerte": rec.get("mehrwerte"),
                        "gruppe_1": rec.get("gruppe_1"),
                        "gruppe_2": rec.get("gruppe_2"),
                        "gesellschafter_field": rec.get("gesellschafter_field"),
                    },
                )

    return stats


def enrich_regions(dry_run: bool = False) -> dict:
    """Populate region + region_prep for records with missing region, using city lookup.

    Reads Städte-Regionen-Matching sheet via region_lookup module (cached).
    Returns stats: total_checked, matched, updated.
    """
    from src.pipeline.region_lookup import load_region_mapping
    from src.pipeline.db import ensure_schema, get_connection

    mapping = load_region_mapping()
    db_path = settings.PIPELINE_DB_PATH
    ensure_schema(db_path)

    with get_connection(db_path) as conn:
        rows = conn.execute(
            "SELECT domain, city FROM company_records "
            "WHERE (region IS NULL OR region = '') AND city IS NOT NULL"
        ).fetchall()

    updates: list[tuple[str, str, str]] = []
    for row in rows:
        city_key = (row["city"] or "").strip().lower()
        entry = mapping.get(city_key)
        if entry is not None:
            region, region_prep = entry
            updates.append((region, region_prep, row["domain"]))

    stats = {"total_checked": len(rows), "matched": len(updates), "updated": 0}

    if dry_run:
        return stats

    with get_connection(db_path) as conn:
        conn.executemany(
            "UPDATE company_records SET region = ?, region_prep = ? WHERE domain = ?",
            updates,
        )
    stats["updated"] = len(updates)
    return stats


def ingest_manual(
    excel_path: Path,
    sheet_name: str = "ALLEX input",
    dry_run: bool = False,
    col_domain: int = 3,
    col_name: int = 4,
    col_klass: int = 7,
    col_source: int = 1,
) -> dict:
    """Ingest pre-classified companies from a manual Excel input file.

    Default columns: Domain (D/3), Source (B/1), Name (E/4), Klass (H/7).
    BA9 format: Domain (E/4), Source (B/1), Name (F/5), Klass (K/10).

    Accepts rows with 'Distributor' (legacy) or any known BA9 category
    code (DEA, INT, SER_PLA_1, etc.).

    New domains → INSERT with source='MANUAL', pipeline_stage='classified'.
    Existing domains → UPDATE source='MANUAL', pipeline_stage='classified'.
    Category codes drive klass + prio + category defaults.
    """
    import hashlib as _hashlib

    from src.pipeline.category_defaults import (
        TARGET_CATEGORIES,
        NON_TARGET_CATEGORIES,
        apply_category_defaults,
        get_default_prio,
    )
    from src.pipeline.db import ensure_schema, get_connection

    _ALL_CATEGORIES = TARGET_CATEGORIES | NON_TARGET_CATEGORIES
    _LEGACY_KLASS = {"A", "B", "C", "D", "E", "S"}
    _ACCEPTED_MANUAL = _ALL_CATEGORIES | _LEGACY_KLASS | {"Distributor"}

    db_path = settings.PIPELINE_DB_PATH

    if not dry_run:
        ensure_schema(db_path)

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(min_row=3, values_only=True))

    stats: dict[str, int] = {
        "total_rows": len(rows),
        "accepted": 0,
        "inserted": 0,
        "updated": 0,
        "skipped_no_domain": 0,
    }

    ingested_at = datetime.now(timezone.utc).isoformat()
    to_insert: list[dict] = []

    for row in rows:
        klass_new = str(row[col_klass] or "").strip() if len(row) > col_klass else ""
        if klass_new not in _ACCEPTED_MANUAL:
            continue
        stats["accepted"] += 1

        domain = normalize_domain(row[col_domain])
        if not domain:
            stats["skipped_no_domain"] += 1
            continue

        name = str(row[col_name] or "").strip() if len(row) > col_name else ""
        source_raw = str(row[col_source] or "").strip() if len(row) > col_source else ""

        # Map category: 'Distributor' → legacy 'B', otherwise use the code directly
        klass_value = "B" if klass_new == "Distributor" else klass_new
        # Prio: new categories get proper prio, legacy 'B' (Distributor) gets 'Prio 1'
        if klass_value in _ALL_CATEGORIES:
            prio_value = get_default_prio(klass_value)
        elif klass_value in ("A", "B"):
            prio_value = "Prio 1"
        elif klass_value == "S":
            prio_value = "Excluded"
        elif klass_value in ("C", "D", "E"):
            prio_value = "Prio 2 (other)"
        else:
            prio_value = None

        rec_id = _hashlib.md5(domain.encode()).hexdigest()[:12]
        rec = {
            "id": rec_id,
            "domain": domain,
            "full_name": name or domain,
            "source": "MANUAL",
            "source_detail": source_raw,
            "klass": klass_value,
            "prio": prio_value,
            "pipeline_stage": "classified",
            "profile_id": "medtech_germany",
            "ingested_at": ingested_at,
            "filter_pass": 1,
            # Category default fields — filled by apply_category_defaults for BA9 codes
            "leistung_text": None,
            "leistung_absatz_2": None,
            "mehrwerte": None,
            "gruppe_1": None,
            "gruppe_2": None,
        }

        # Apply category defaults for BA9 codes
        if klass_value in _ALL_CATEGORIES:
            apply_category_defaults(rec, klass_value)

        to_insert.append(rec)

    if dry_run:
        logger.info(
            "DRY RUN — would ingest %d records (%d skipped no domain)",
            len(to_insert),
            stats["skipped_no_domain"],
        )
        return stats

    with get_connection(db_path) as conn:
        row = conn.execute(
            "SELECT MAX(CAST(SUBSTR(briefaktion, 3) AS INTEGER)) FROM company_records WHERE briefaktion LIKE 'BA%'"
        ).fetchone()
        next_ba = f"BA{(row[0] or 0) + 1}"
        logger.info("Assigning briefaktion %s to %d records", next_ba, len(to_insert))
        stats["briefaktion"] = next_ba

        for rec in to_insert:
            existing = conn.execute(
                "SELECT domain FROM company_records WHERE domain = ?",
                (rec["domain"],),
            ).fetchone()

            rec["briefaktion"] = next_ba

            if existing:
                conn.execute(
                    """UPDATE company_records
                       SET source = 'MANUAL',
                           klass = :klass,
                           prio = COALESCE(:prio, prio),
                           pipeline_stage = 'classified',
                           filter_pass = 1,
                           briefaktion = :briefaktion,
                           is_subsidiary = NULL,
                           is_pe_backed = NULL,
                           reclassify_reason = NULL,
                           full_name = COALESCE(full_name, :full_name),
                           leistung_text = COALESCE(:leistung_text, leistung_text),
                           leistung_absatz_2 = COALESCE(:leistung_absatz_2, leistung_absatz_2),
                           mehrwerte = COALESCE(:mehrwerte, mehrwerte),
                           gruppe_1 = COALESCE(:gruppe_1, gruppe_1),
                           gruppe_2 = COALESCE(:gruppe_2, gruppe_2)
                     WHERE domain = :domain""",
                    rec,
                )
                stats["updated"] += 1
                logger.info(
                    "Updated existing: %s → MANUAL/%s (%s)",
                    rec["domain"],
                    rec["klass"],
                    next_ba,
                )
            else:
                conn.execute(
                    """INSERT INTO company_records
                       (id, domain, full_name, profile_id, source, pipeline_stage,
                        klass, prio, filter_pass, ingested_at, briefaktion,
                        leistung_text, leistung_absatz_2, mehrwerte, gruppe_1, gruppe_2)
                     VALUES
                       (:id, :domain, :full_name, :profile_id, :source, :pipeline_stage,
                        :klass, :prio, :filter_pass, :ingested_at, :briefaktion,
                        :leistung_text, :leistung_absatz_2, :mehrwerte, :gruppe_1, :gruppe_2)""",
                    rec,
                )
                stats["inserted"] += 1
                logger.info("Inserted new: %s (%s)", rec["domain"], rec["full_name"])

    return stats
