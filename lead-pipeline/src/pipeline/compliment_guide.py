"""Build a compliment style guide from the Serienbriefe corpus.

Reads Kompliment 1 and Kompliment 2 from the source Excel, cross-references
pipeline.db for positive-response weighting, then calls Claude once to synthesise
a style guide saved to src/config/compliment_guide.md.

Usage:
    python pipeline.py build-compliment-guide [--via-cli] [--dry-run]
"""

from __future__ import annotations

import json
import logging
import os
import subprocess
from pathlib import Path
from typing import Optional

import openpyxl

from src.config import settings

logger = logging.getLogger(__name__)

_GUIDE_OUTPUT_PATH = Path(__file__).parent.parent / "config" / "compliment_guide.md"

# Positive outreach statuses (weight these examples higher in the guide)
_POSITIVE_STATUSES = {"contact", "meeting", "financials", "offer", "deal"}

_GUIDE_PROMPT = """You are a German business writing expert. Analyse these Kompliment sentences from M&A outreach letters sent to German Medizintechnik distribution companies.

Kompliment 1 appears in the opening paragraph of the letter. It references the company's history, longevity, or accumulated expertise.
Kompliment 2 appears in the second paragraph. It highlights a specific service concept, product specialisation, or differentiator.

Below are examples from letters that received POSITIVE responses (meetings, contact made, financials shared) — these are the most effective ones:

=== POSITIVE-RESPONSE EXAMPLES ===
{positive_examples}

=== ALL OTHER EXAMPLES (sample) ===
{other_examples}

Please write a concise style guide in Markdown covering:
1. **Kompliment 1** — position in letter, grammatical pattern, typical length, most common opening words, 5 best examples (from positive-response set first), 3 patterns to avoid
2. **Kompliment 2** — same structure as above
3. **General rules** — what distinguishes effective compliments from generic ones in this corpus

Be direct and specific. The guide will be used by an AI to generate new compliments for similar companies.
"""


def _load_corpus(excel_path: Path) -> list[dict]:
    """Read K1/K2/domain from Serienbriefe sheet. Returns list of dicts."""
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb[settings.SERIENBRIEFE_SHEET]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    corpus = []
    for row in rows:
        if len(row) <= 23:
            continue
        domain = str(row[2]).strip().lower() if row[2] else None
        k1 = str(row[22]).strip() if row[22] else None
        k2 = str(row[23]).strip() if row[23] else None
        if domain and (k1 or k2):
            corpus.append({"domain": domain, "k1": k1, "k2": k2})
    return corpus


def _get_positive_domains(db_path: Path) -> set[str]:
    """Return domains with positive outreach_status from pipeline.db."""
    import sqlite3

    if not db_path.exists():
        return set()
    conn = sqlite3.connect(db_path)
    placeholders = ",".join("?" * len(_POSITIVE_STATUSES))
    rows = conn.execute(
        f"SELECT domain FROM company_records WHERE outreach_status IN ({placeholders})",
        list(_POSITIVE_STATUSES),
    ).fetchall()
    conn.close()
    return {r[0] for r in rows}


def _format_examples(entries: list[dict], max_count: int = 30) -> str:
    lines = []
    for e in entries[:max_count]:
        if e.get("k1"):
            lines.append(f"K1: {e['k1']}")
        if e.get("k2"):
            lines.append(f"K2: {e['k2']}")
        lines.append("")
    return "\n".join(lines).strip()


def _call_claude_cli(prompt: str) -> Optional[str]:
    """Call Claude CLI subprocess to generate the guide."""
    env = os.environ.copy()
    env.pop("ANTHROPIC_API_KEY", None)
    try:
        result = subprocess.run(
            [settings.CLAUDE_CMD, "-p", "--output-format", "text"],
            input=prompt,
            capture_output=True,
            text=True,
            timeout=120,
            encoding="utf-8",
            env=env,
        )
        if result.returncode != 0:
            logger.error("Claude CLI error: %s", result.stderr.strip())
            return None
        return result.stdout.strip()
    except Exception as e:
        logger.error("Claude CLI call failed: %s", e)
        return None


def build_guide(
    dry_run: bool = False,
    db_path: Optional[Path] = None,
    output_path: Optional[Path] = None,
) -> dict:
    """Build compliment_guide.md from the Serienbriefe corpus.

    Returns stats dict: {k1_count, k2_count, positive_count, guide_written}.
    """
    if db_path is None:
        db_path = settings.PIPELINE_DB_PATH
    if output_path is None:
        output_path = _GUIDE_OUTPUT_PATH

    excel_path = settings.SOURCE_EXCEL
    if not excel_path.exists():
        raise FileNotFoundError(f"Source Excel not found: {excel_path}")

    corpus = _load_corpus(excel_path)
    positive_domains = _get_positive_domains(db_path)

    positive = [e for e in corpus if e["domain"] in positive_domains]
    other = [e for e in corpus if e["domain"] not in positive_domains]

    k1_count = sum(1 for e in corpus if e.get("k1"))
    k2_count = sum(1 for e in corpus if e.get("k2"))

    stats = {
        "k1_count": k1_count,
        "k2_count": k2_count,
        "positive_count": len(positive),
        "total_count": len(corpus),
        "guide_written": False,
    }

    logger.info(
        "Corpus: %d entries — K1=%d K2=%d positive=%d",
        len(corpus),
        k1_count,
        k2_count,
        len(positive),
    )

    if dry_run:
        logger.info("DRY RUN — would call Claude once; guide NOT written")
        return stats

    prompt = _GUIDE_PROMPT.format(
        positive_examples=_format_examples(positive, max_count=60),
        other_examples=_format_examples(other, max_count=30),
    )

    logger.info(
        "Calling Claude to synthesise compliment guide (%d chars prompt)...",
        len(prompt),
    )
    guide_text = _call_claude_cli(prompt)

    if not guide_text:
        logger.error("Failed to generate guide — no output from Claude")
        return stats

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(guide_text, encoding="utf-8")
    logger.info("Guide written to %s (%d chars)", output_path, len(guide_text))
    stats["guide_written"] = True
    return stats
